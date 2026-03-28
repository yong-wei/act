#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


QUESTION_HEADING_RE = re.compile(r'^(?P<number>\d+-\d+)(?P<suffix>.*)$')
BODY_TEXT_QUESTION_RE = re.compile(r'^(?P<number>\d+-\d+)\s+(?P<rest>.+)$')
SOLUTION_START_RE = re.compile(r'^(?:解|证明)[:：\s　]')
FIGURE_TITLE_RE = re.compile(r'^图\s*\d')
FIGURE_PROMPT_RE = re.compile(r'(试|求|要求|确定|说明|分析|讨论|绘制|画出|求出|试问|计算|设计)')
POINT_RE = re.compile(r'[\(（](\d+)分[\)）]')
BLOCK_LATEX_RE = re.compile(r'^\[\s*(.+?)\s*\]$')
INLINE_FORMULA_CANDIDATE_RE = re.compile(r'\(\s*([^()\n]{1,200})\s*\)')
LATEX_MARKER_RE = re.compile(r'(\\[A-Za-z]+|[_^]|\\frac|\\omega|\\theta|\\sigma|\\zeta|\\epsilon|\\Omega)')
MAX_BODY_FALLBACK_CHAPTER = 8
CJK_RE = re.compile(r'[\u4e00-\u9fff]')
FORMULA_SYMBOL_RE = re.compile(r'(=|\\cdot|\\lim|\\sum|\\prod|\\int|\\to|/|\+|-|\*|[A-Za-z]+\([A-Za-z0-9])')
UNRESOLVED_BRACKET_FORMULA_RE = re.compile(r'(?<!\\left)\[[^\]\n]*\\[A-Za-z]+[^\]\n]*(?<!\\right)\]')


@dataclass
class QuestionSegment:
    index: int
    question_id: str
    question_number: str
    source_ref: str
    chapter: int
    section: str | None
    paragraphs: list[Paragraph]
    next_paragraph: Paragraph | None


def parse_args() -> argparse.Namespace:
    default_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description='Extract structured questions from DOCX question bank.')
    parser.add_argument(
        '--source',
        default=str(default_root / 'source' / '自动控制原理习题解析.docx'),
        help='Path to source DOCX.',
    )
    parser.add_argument(
        '--output-root',
        default=str(default_root),
        help='Output root containing questions/assets/indexes/reports directories.',
    )
    parser.add_argument('--limit', type=int, default=None, help='Only extract first N questions for testing.')
    return parser.parse_args()


def normalize_text(text: str) -> str:
    lines: list[str] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        block_match = BLOCK_LATEX_RE.match(line)
        if block_match and LATEX_MARKER_RE.search(block_match.group(1)):
            lines.append(f'$$ {clean_formula(block_match.group(1))} $$')
            continue
        line = normalize_inline_formulas(line)
        line = normalize_enclosed_formulas(line, '[', ']')
        line = normalize_enclosed_formulas(line, '(', ')')
        lines.append(line)
    return '\n\n'.join(lines).strip()


def clean_formula(content: str) -> str:
    return re.sub(r'\s+', ' ', content.strip())


def normalize_inline_formulas(line: str) -> str:
    def replace(match: re.Match[str]) -> str:
        content = match.group(1).strip()
        if not LATEX_MARKER_RE.search(content):
            return match.group(0)
        if not has_formula_safe_boundaries(line, match.start(), match.end()):
            return match.group(0)
        return f'${clean_formula(content)}$'

    return INLINE_FORMULA_CANDIDATE_RE.sub(replace, line)


def looks_like_formula(content: str) -> bool:
    compact = content.strip()
    if not compact:
        return False
    if LATEX_MARKER_RE.search(compact):
        return True
    if CJK_RE.search(compact):
        return False
    return bool(FORMULA_SYMBOL_RE.search(compact))


def has_formula_safe_boundaries(line: str, start: int, end: int) -> bool:
    left_index = start - 1
    while left_index >= 0 and line[left_index].isspace():
        left_index -= 1

    right_index = end
    while right_index < len(line) and line[right_index].isspace():
        right_index += 1

    left_char = line[left_index] if left_index >= 0 else ''
    right_char = line[right_index] if right_index < len(line) else ''
    unsafe_chars = set('\\${}[]_^') | set('abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789')

    return left_char not in unsafe_chars and right_char not in unsafe_chars


def normalize_enclosed_formulas(line: str, opening: str, closing: str) -> str:
    result: list[str] = []
    index = 0
    while index < len(line):
        if line[index] != opening:
            result.append(line[index])
            index += 1
            continue

        depth = 1
        cursor = index + 1
        while cursor < len(line) and depth > 0:
            if line[cursor] == opening:
                depth += 1
            elif line[cursor] == closing:
                depth -= 1
            cursor += 1

        if depth != 0:
            result.append(line[index])
            index += 1
            continue

        content = line[index + 1:cursor - 1].strip()
        if looks_like_formula(content) and (
            opening != '(' or has_formula_safe_boundaries(line, index, cursor)
        ):
            prefix = ''.join(result).strip()
            suffix = line[cursor:].strip()
            formula = clean_formula(content)
            result.append(f'$$ {formula} $$' if not prefix and not suffix else f'${formula}$')
        else:
            result.append(line[index:cursor])
        index = cursor

    return ''.join(result)


def extract_images_from_element(element: object, part: object) -> list[tuple[bytes, str]]:
    images: list[tuple[bytes, str]] = []
    seen_rids: set[str] = set()
    for node in element.xpath('.//*[local-name()="blip"]'):
        rid = node.get(qn('r:embed'))
        if not rid or rid in seen_rids:
            continue
        seen_rids.add(rid)
        image_part = part.related_parts[rid]
        suffix = Path(str(image_part.partname)).suffix or '.bin'
        images.append((image_part.blob, suffix))
    return images


def has_malformed_dollar_delimiters(text: str) -> bool:
    active: str | None = None
    index = 0

    while index < len(text):
        if text[index] != '$':
            index += 1
            continue

        token = '$$' if text[index:index + 2] == '$$' else '$'
        index += len(token)

        if active is None:
            active = token
            continue

        if active == token:
            active = None
            continue

        return True

    return active is not None


def classify_formula_status(stem_md: str, solution_md: str) -> str:
    combined = f'{stem_md}\n{solution_md}'
    unresolved = has_malformed_dollar_delimiters(combined) or UNRESOLVED_BRACKET_FORMULA_RE.search(combined)
    return 'mixed' if unresolved else 'clean'


def derive_tags(section: str | None, stem_md: str, solution_md: str) -> list[str]:
    text = f'{section or ""}\n{stem_md}\n{solution_md}'
    tags: list[str] = []
    keyword_map = {
        '方框图': '方框图',
        '信号流图': '信号流图',
        '微分方程': '微分方程',
        '传递函数': '传递函数',
        '稳定': '稳定性',
        '根轨迹': '根轨迹',
        '频率': '频域分析',
        'Bode': '频域分析',
        'Nyquist': '频域分析',
        '状态空间': '状态空间',
        '非线性': '非线性系统',
        '液位': '液位控制',
        '温度': '温度控制',
        '机械系统': '机械系统',
    }
    if section:
        tags.append(section)
    for keyword, tag in keyword_map.items():
        if keyword in text and tag not in tags:
            tags.append(tag)
    return tags


def extract_question_number(text: str, style_name: str, allow_body_fallback: bool = False) -> str | None:
    if style_name.startswith('Heading'):
        match = QUESTION_HEADING_RE.match(text)
        return match.group('number') if match else None

    if not allow_body_fallback:
        return None

    match = BODY_TEXT_QUESTION_RE.match(text)
    if not match:
        return None

    number = match.group('number')
    chapter_text, _ = number.split('-', 1)
    try:
        chapter = int(chapter_text)
    except ValueError:
        return None

    rest = match.group('rest').lstrip()
    if not (1 <= chapter <= MAX_BODY_FALLBACK_CHAPTER):
        return None
    if not rest or rest[0] in '-/|':
        return None

    return number


def build_segments(document: Document, limit: int | None = None) -> list[QuestionSegment]:
    heading_question_indices: list[tuple[int, str, str | None]] = []
    body_fallback_indices: list[tuple[int, str, str | None]] = []
    current_section: str | None = None

    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        style_name = paragraph.style.name if paragraph.style else ''
        if not text:
            continue

        question_number = extract_question_number(text, style_name)
        if style_name.startswith('Heading') and not question_number:
            current_section = text
            continue

        if question_number:
            heading_question_indices.append((index, question_number, current_section))
            continue

        question_number = extract_question_number(text, style_name, allow_body_fallback=True)
        if question_number:
            body_fallback_indices.append((index, question_number, current_section))

    heading_numbers = {question_number for _, question_number, _ in heading_question_indices}
    question_indices = heading_question_indices + [
        item for item in body_fallback_indices if item[1] not in heading_numbers
    ]
    question_indices.sort(key=lambda item: item[0])

    if limit is not None:
        question_indices = question_indices[:limit]

    segments: list[QuestionSegment] = []
    for offset, (start_index, question_number, section) in enumerate(question_indices):
        end_index = (
            question_indices[offset + 1][0]
            if offset + 1 < len(question_indices)
            else len(document.paragraphs)
        )
        question_id = f'AC-Q-{offset + 1:04d}'
        paragraphs = document.paragraphs[start_index:end_index]
        segments.append(
            QuestionSegment(
                index=offset + 1,
                question_id=question_id,
                question_number=question_number,
                source_ref=paragraphs[0].text.strip(),
                chapter=int(question_number.split('-', 1)[0]),
                section=section,
                paragraphs=paragraphs,
                next_paragraph=document.paragraphs[end_index] if end_index < len(document.paragraphs) else None,
            )
        )
    return segments


def classify_figure_line(text: str) -> str | None:
    if not FIGURE_TITLE_RE.match(text):
        return None
    if FIGURE_PROMPT_RE.search(text):
        return 'reference'
    return 'title'


def split_stem_and_solution(paragraphs: list[Paragraph]) -> tuple[list[str], list[str], list[str], list[str]]:
    stem_lines: list[str] = []
    solution_lines: list[str] = []
    figure_titles: list[str] = []
    figure_references: list[str] = []
    in_solution = False

    for paragraph in paragraphs[1:]:
        text = paragraph.text.strip()
        if not text:
            continue
        figure_line_kind = classify_figure_line(text)
        if figure_line_kind == 'title':
            figure_titles.append(text)
        elif figure_line_kind == 'reference':
            figure_references.append(text)
        if SOLUTION_START_RE.match(text):
            in_solution = True

        target = solution_lines if in_solution else stem_lines
        target.append(text)

    return stem_lines, solution_lines, figure_titles, figure_references


def build_inline_score_points(solution_md: str) -> list[dict[str, object]]:
    points: list[dict[str, object]] = []
    for line in solution_md.splitlines():
        for match in POINT_RE.finditer(line):
            points.append({
                'points': int(match.group(1)),
                'excerpt': line.strip(),
            })
    return points


def build_rubric(inline_score_points: list[dict[str, object]]) -> list[dict[str, object]]:
    rubric: list[dict[str, object]] = []
    for index, item in enumerate(inline_score_points, start=1):
        rubric.append({
            'step': f'步骤 {index}',
            'points': item['points'],
            'criteria': item['excerpt'],
        })
    return rubric


def write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def load_curated_payloads(output_root: Path) -> dict[tuple[str, str], dict[str, object]]:
    curated: dict[tuple[str, str], dict[str, object]] = {}
    questions_dir = output_root / 'questions'
    if not questions_dir.exists():
        return curated

    for path in questions_dir.glob('AC-Q-*.json'):
        try:
            payload = json.loads(path.read_text(encoding='utf-8'))
        except (OSError, json.JSONDecodeError):
            continue

        usage_status = str(payload.get('usage_status', 'raw'))
        question_number = str(payload.get('question_number', ''))
        source_ref = str(payload.get('source_ref', ''))
        if usage_status == 'raw' or not question_number or not source_ref:
            continue

        curated[(question_number, source_ref)] = payload

    return curated


def merge_curated_payload(
    payload: dict[str, object],
    curated_payload: dict[str, object] | None,
) -> dict[str, object]:
    if not curated_payload:
        return payload

    merged = dict(payload)
    for field in [
        'stem_md',
        'solution_md',
        'inline_score_points',
        'rubric',
        'formula_status',
        'usage_status',
    ]:
        if field in curated_payload:
            merged[field] = curated_payload[field]

    if curated_payload.get('knowledge_tags'):
        merged['knowledge_tags'] = curated_payload['knowledge_tags']

    return merged


def dump_question_markdown(payload: dict[str, object], question_root: Path) -> str:
    assets_md = '\n'.join(
        f'- `{title}` -> `../{asset}`'
        for title, asset in zip(payload['figure_titles'], payload['figure_assets'])
    )
    if not assets_md and payload['figure_assets']:
        assets_md = '\n'.join(f'- `../{asset}`' for asset in payload['figure_assets'])
    if not assets_md and payload['figure_titles']:
        assets_md = '\n'.join(f'- 图题：`{title}`' for title in payload['figure_titles'])
    if not assets_md and payload.get('figure_references'):
        assets_md = '\n'.join(f'- 题面引用：`{title}`' for title in payload['figure_references'])
    if not assets_md:
        assets_md = '- 暂无'

    points_md = '\n'.join(
        f"- {item['points']} 分：{item['excerpt']}"
        for item in payload['inline_score_points']
    ) or '- 暂无'

    rubric_md = '\n'.join(
        f"- {item['step']} | {item['points']} 分：{item['criteria']}"
        for item in payload['rubric']
    ) or '- 暂无'

    section = payload['section'] or '未标注'
    return f"""---
question_id: {payload['question_id']}
source_ref: {payload['source_ref']}
question_number: {payload['question_number']}
chapter: {payload['chapter']}
section: {section}
figure_status: {payload['figure_status']}
formula_status: {payload['formula_status']}
usage_status: {payload['usage_status']}
---

## 题面

{payload['stem_md']}

### 配图

{assets_md}

## 答案解析

{payload['solution_md']}

## 行内得分点

{points_md}

## 评分指南

{rubric_md}

## 元数据

- 来源：`{payload['source_ref']}`
- 章节：`{payload['chapter']}`
- 小节：`{section}`
- 标签：{', '.join(payload['knowledge_tags']) if payload['knowledge_tags'] else '暂无'}
"""


def iter_segment_block_elements(segment: QuestionSegment) -> Iterable[object]:
    current = segment.paragraphs[0]._p
    end = segment.next_paragraph._p if segment.next_paragraph is not None else None

    while current is not None and current is not end:
        yield current
        current = current.getnext()


def export_question_images(
    segment: QuestionSegment,
    output_root: Path,
) -> list[str]:
    asset_dir = output_root / 'assets' / segment.question_id
    if asset_dir.exists():
        shutil.rmtree(asset_dir)
    asset_dir.mkdir(parents=True, exist_ok=True)

    exported: list[str] = []
    seen_bytes: set[bytes] = set()
    counter = 1
    part = segment.paragraphs[0].part
    for element in iter_segment_block_elements(segment):
        for blob, suffix in extract_images_from_element(element, part):
            if blob in seen_bytes:
                continue
            seen_bytes.add(blob)
            filename = f'image-{counter:02d}{suffix}'
            target = asset_dir / filename
            target.write_bytes(blob)
            exported.append(str(target.relative_to(output_root)).replace('\\', '/'))
            counter += 1
    if not exported:
        shutil.rmtree(asset_dir)
    return exported


def extract_question_payload(segment: QuestionSegment, output_root: Path) -> dict[str, object]:
    stem_lines, solution_lines, figure_titles, figure_references = split_stem_and_solution(segment.paragraphs)
    stem_md = normalize_text('\n'.join(stem_lines))
    solution_md = normalize_text('\n'.join(solution_lines))
    figure_assets = export_question_images(segment, output_root)
    inline_score_points = build_inline_score_points(solution_md)
    rubric = build_rubric(inline_score_points)

    if figure_assets:
        figure_status = 'complete'
    elif figure_titles:
        figure_status = 'source_missing'
    elif figure_references:
        figure_status = 'text_reference_only'
    else:
        figure_status = 'none'

    payload: dict[str, object] = {
        'question_id': segment.question_id,
        'source_ref': segment.source_ref,
        'question_number': segment.question_number,
        'chapter': segment.chapter,
        'section': segment.section,
        'stem_md': stem_md,
        'solution_md': solution_md,
        'inline_score_points': inline_score_points,
        'rubric': rubric,
        'figure_titles': figure_titles,
        'figure_references': figure_references,
        'figure_assets': figure_assets,
        'figure_status': figure_status,
        'formula_status': classify_formula_status(stem_md, solution_md),
        'usage_status': 'raw',
        'knowledge_tags': derive_tags(segment.section, stem_md, solution_md),
    }
    return payload


def write_reports(output_root: Path, payloads: Iterable[dict[str, object]]) -> None:
    payload_list = list(payloads)
    unresolved = [
        {
            'question_id': item['question_id'],
            'source_ref': item['source_ref'],
            'formula_status': item['formula_status'],
        }
        for item in payload_list
        if item['formula_status'] != 'clean'
    ]
    missing_images = [
        {
            'question_id': item['question_id'],
            'source_ref': item['source_ref'],
            'figure_status': item['figure_status'],
            'figure_titles': item['figure_titles'],
            'figure_references': item.get('figure_references', []),
        }
        for item in payload_list
        if item['figure_status'] in {'source_missing', 'text_reference_only'}
    ]
    figure_status_breakdown: dict[str, int] = {}
    for item in payload_list:
        status = str(item['figure_status'])
        figure_status_breakdown[status] = figure_status_breakdown.get(status, 0) + 1
    source_missing_count = figure_status_breakdown.get('source_missing', 0)
    text_reference_only_count = figure_status_breakdown.get('text_reference_only', 0)
    missing_figure_count = source_missing_count + text_reference_only_count

    source_refs: dict[str, list[str]] = {}
    for item in payload_list:
        source_refs.setdefault(str(item['source_ref']), []).append(str(item['question_id']))
    duplicates = [
        {'source_ref': key, 'question_ids': ids}
        for key, ids in source_refs.items()
        if len(ids) > 1
    ]

    write_json(output_root / 'reports' / 'unresolved-formulas.json', unresolved)
    write_json(output_root / 'reports' / 'missing-images.json', missing_images)
    write_json(output_root / 'reports' / 'duplicate-candidates.json', duplicates)
    write_json(
        output_root / 'reports' / 'extraction-report.json',
        {
            'question_count': len(payload_list),
            # Keep the legacy key for existing consumers while exposing the new status-aware name.
            'title_only_figures': missing_figure_count,
            'missing_figure_questions': missing_figure_count,
            'source_missing_questions': source_missing_count,
            'text_reference_only_questions': text_reference_only_count,
            'figure_status_breakdown': figure_status_breakdown,
            'mixed_formula_questions': len(unresolved),
            'duplicate_candidates': len(duplicates),
        },
    )


def main() -> None:
    args = parse_args()
    source = Path(args.source)
    output_root = Path(args.output_root)
    output_root.mkdir(parents=True, exist_ok=True)
    for name in ['questions', 'assets', 'reports']:
        (output_root / name).mkdir(parents=True, exist_ok=True)

    document = Document(source)
    segments = build_segments(document, args.limit)
    curated_payloads = load_curated_payloads(output_root)
    payloads: list[dict[str, object]] = []

    for segment in segments:
        payload = extract_question_payload(segment, output_root)
        payload = merge_curated_payload(
            payload,
            curated_payloads.get((str(payload['question_number']), str(payload['source_ref']))),
        )
        question_json_path = output_root / 'questions' / f'{segment.question_id}.json'
        question_md_path = output_root / 'questions' / f'{segment.question_id}.md'
        write_json(question_json_path, payload)
        question_md_path.write_text(dump_question_markdown(payload, output_root), encoding='utf-8')
        payloads.append(payload)

    write_reports(output_root, payloads)
    print(json.dumps({'question_count': len(payloads)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
