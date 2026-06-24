#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


FONT_HEITI = 'SimHei'
FONT_SONGTI = 'SimSun'
FONT_WESTERN = 'Times New Roman'
COLOR_BLACK = RGBColor(0, 0, 0)

SIZE_H1 = 16
SIZE_H2 = 14
SIZE_BODY = 12
SIZE_CAPTION = 10.5
BODY_FIRST_LINE_INDENT = 24


def apply_run_font(run, east_asia: str, size: float, *, bold: bool = False) -> None:
    run.font.name = FONT_WESTERN
    run._element.rPr.rFonts.set(qn('w:ascii'), FONT_WESTERN)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), FONT_WESTERN)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), east_asia)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = COLOR_BLACK


def set_run_font(run, size: int, bold: bool = False, math: bool = False) -> None:
    apply_run_font(run, FONT_SONGTI, size, bold=bold)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(8 if level > 1 else 0)
    paragraph.paragraph_format.space_after = Pt(6)
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        east_asia = FONT_HEITI
        size = SIZE_H1
        bold = False
    elif level == 2:
        east_asia = FONT_SONGTI
        size = SIZE_H2
        bold = True
    elif level == 3:
        east_asia = FONT_HEITI
        size = SIZE_BODY
        bold = False
    else:
        east_asia = FONT_SONGTI
        size = SIZE_BODY
        bold = True
    run = paragraph.add_run(text)
    apply_run_font(run, east_asia, size, bold=bold)


def add_paragraph(doc: Document, text: str, *, style: str | None = None, center: bool = False, math: bool = False) -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.first_line_indent = Pt(BODY_FIRST_LINE_INDENT)
    paragraph.paragraph_format.space_after = Pt(4)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, SIZE_BODY, bold=False, math=math)


def add_image(doc: Document, image_path: Path, alt_text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(6.0))
    if alt_text:
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.line_spacing = 1.0
        caption.paragraph_format.first_line_indent = Pt(0)
        caption.paragraph_format.space_after = Pt(6)
        caption_run = caption.add_run(alt_text)
        apply_run_font(caption_run, FONT_SONGTI, SIZE_CAPTION)


def style_name(paragraph) -> str:
    return (paragraph.style.name or '').lower() if paragraph.style is not None else ''


def paragraph_has_drawing(paragraph) -> bool:
    return bool(paragraph._element.xpath('.//w:drawing'))


def previous_paragraph_has_drawing(doc: Document, index: int) -> bool:
    if index <= 0:
        return False
    return paragraph_has_drawing(doc.paragraphs[index - 1])


def is_heading(paragraph, level: int) -> bool:
    name = style_name(paragraph)
    return name in {f'heading {level}', f'标题 {level}'}


def is_caption(doc: Document, index: int, paragraph) -> bool:
    name = style_name(paragraph)
    if 'caption' in name or '题注' in name:
        return True
    if paragraph.text.strip() and previous_paragraph_has_drawing(doc, index):
        return True
    return False


def apply_paragraph_spacing(paragraph, *, first_line_indent: bool) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.0
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(4)
    fmt.first_line_indent = Pt(BODY_FIRST_LINE_INDENT if first_line_indent else 0)


def format_paragraph(doc: Document, index: int, paragraph) -> None:
    if is_heading(paragraph, 1):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        apply_paragraph_spacing(paragraph, first_line_indent=False)
        for run in paragraph.runs:
            apply_run_font(run, FONT_HEITI, SIZE_H1)
        return

    if is_heading(paragraph, 2):
        apply_paragraph_spacing(paragraph, first_line_indent=False)
        for run in paragraph.runs:
            apply_run_font(run, FONT_SONGTI, SIZE_H2, bold=True)
        return

    if is_heading(paragraph, 3):
        apply_paragraph_spacing(paragraph, first_line_indent=False)
        for run in paragraph.runs:
            apply_run_font(run, FONT_HEITI, SIZE_BODY)
        return

    if is_caption(doc, index, paragraph):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        apply_paragraph_spacing(paragraph, first_line_indent=False)
        for run in paragraph.runs:
            apply_run_font(run, FONT_SONGTI, SIZE_CAPTION)
        return

    apply_paragraph_spacing(paragraph, first_line_indent=bool(paragraph.text.strip()))
    for run in paragraph.runs:
        if run.bold or run.italic:
            apply_run_font(run, FONT_HEITI, SIZE_BODY)
            run.italic = False
        else:
            apply_run_font(run, FONT_SONGTI, SIZE_BODY)


def format_table(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                apply_paragraph_spacing(paragraph, first_line_indent=False)
                for run in paragraph.runs:
                    apply_run_font(run, FONT_SONGTI, SIZE_CAPTION)


def apply_homework_docx_format(output_path: Path) -> None:
    doc = Document(str(output_path))
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.27)
        section.bottom_margin = Cm(1.27)
        section.left_margin = Cm(1.27)
        section.right_margin = Cm(1.27)

    for style in doc.styles:
        if getattr(style, 'font', None) is not None:
            style.font.name = FONT_WESTERN
            style.element.rPr.rFonts.set(qn('w:ascii'), FONT_WESTERN)
            style.element.rPr.rFonts.set(qn('w:hAnsi'), FONT_WESTERN)
            style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_SONGTI)
            style.font.color.rgb = COLOR_BLACK

    for index, paragraph in enumerate(doc.paragraphs):
        format_paragraph(doc, index, paragraph)

    for table in doc.tables:
        format_table(table)

    doc.save(output_path)


def render_with_pandoc(markdown: str, output_path: Path, base_dir: Path) -> bool:
    pandoc = shutil.which('pandoc')
    if pandoc is None:
        return False

    output_path.parent.mkdir(parents=True, exist_ok=True)
    command = [
        pandoc,
        '--from',
        'markdown+tex_math_dollars+pipe_tables+lists_without_preceding_blankline',
        '--to',
        'docx',
        '--resource-path',
        str(base_dir),
        '--output',
        str(output_path),
    ]
    try:
        subprocess.run(
            command,
            input=markdown,
            text=True,
            check=True,
            capture_output=True,
        )
    except subprocess.CalledProcessError as exc:
        if exc.stderr:
            print(exc.stderr.strip(), file=sys.stderr)
        return False
    apply_homework_docx_format(output_path)
    return True


def flush_paragraph(doc: Document, buffer: list[str]) -> None:
    if not buffer:
        return
    text = ' '.join(part.strip() for part in buffer if part.strip())
    if text:
        add_paragraph(doc, text)
    buffer.clear()


def render_markdown_simple(markdown: str, output_path: Path, base_dir: Path) -> None:
    doc = Document()

    paragraph_buffer: list[str] = []
    math_buffer: list[str] = []
    in_math = False

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped == '$$':
            flush_paragraph(doc, paragraph_buffer)
            if in_math:
                for math_line in math_buffer:
                    add_paragraph(doc, math_line, center=True, math=True)
                math_buffer.clear()
                in_math = False
            else:
                in_math = True
            continue

        if in_math:
            if stripped:
                math_buffer.append(stripped)
            continue

        if not stripped:
            flush_paragraph(doc, paragraph_buffer)
            continue

        if stripped == '---':
            flush_paragraph(doc, paragraph_buffer)
            doc.add_paragraph()
            continue

        heading_match = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if heading_match:
            flush_paragraph(doc, paragraph_buffer)
            level = len(heading_match.group(1))
            add_heading(doc, heading_match.group(2), level)
            continue

        image_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if image_match:
            flush_paragraph(doc, paragraph_buffer)
            alt_text = image_match.group(1) or Path(image_match.group(2)).name
            image_path = (base_dir / image_match.group(2)).resolve()
            if image_path.is_file():
                add_image(doc, image_path, alt_text)
            else:
                add_paragraph(doc, f'[图片缺失] {alt_text}: {image_match.group(2)}')
            continue

        bullet_match = re.match(r'^-\s+(.*)$', stripped)
        if bullet_match:
            flush_paragraph(doc, paragraph_buffer)
            add_paragraph(doc, bullet_match.group(1), style='List Bullet')
            continue

        numbered_match = re.match(r'^\d+\.\s+(.*)$', stripped)
        if numbered_match:
            flush_paragraph(doc, paragraph_buffer)
            add_paragraph(doc, numbered_match.group(1), style='List Number')
            continue

        paragraph_buffer.append(stripped)

    flush_paragraph(doc, paragraph_buffer)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    apply_homework_docx_format(output_path)


def render_markdown(markdown: str, output_path: Path, base_dir: Path) -> None:
    if render_with_pandoc(markdown, output_path, base_dir):
        return
    print(
        'warning: pandoc is unavailable or failed; falling back to simple DOCX export '
        'without native Word math/table conversion',
        file=sys.stderr,
    )
    render_markdown_simple(markdown, output_path, base_dir)


def main() -> int:
    parser = argparse.ArgumentParser(description='Export homework markdown to a simple DOCX file.')
    parser.add_argument('markdown_path', help='Input markdown path')
    parser.add_argument('output_path', nargs='?', help='Optional output .docx path')
    args = parser.parse_args()

    markdown_path = Path(args.markdown_path).resolve()
    output_path = Path(args.output_path).resolve() if args.output_path else markdown_path.with_suffix('.docx')
    render_markdown(markdown_path.read_text(encoding='utf-8'), output_path, markdown_path.parent)
    print(output_path)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
